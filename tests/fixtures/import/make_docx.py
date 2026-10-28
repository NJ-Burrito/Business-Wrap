"""Writes the DOCX import fixture (tests/library.test.mjs): 3 headings, 2 paragraphs, 1 table, Arabic.

    uv run --no-project --with python-docx python tests/fixtures/import/make_docx.py <out.docx>
"""
import sys

from docx import Document

d = Document()
d.add_heading("المقدمة", 1)
d.add_paragraph("هذه وثيقة تجريبية للاستيراد.")
d.add_heading("الغرض من الوثيقة", 2)
d.add_paragraph("يصف هذا القسم الغرض من الخدمة.")
d.add_heading("المتطلبات الوظيفية", 1)
t = d.add_table(rows=3, cols=2)
for i, row in enumerate([["الرمز", "المتطلب"], ["FR-01", "تقديم الطلب"], ["FR-02", "دفع الرسوم"]]):
    for j, v in enumerate(row):
        t.cell(i, j).text = v
d.save(sys.argv[1])
