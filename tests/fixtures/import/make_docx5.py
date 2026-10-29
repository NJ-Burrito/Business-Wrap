"""Writes the structured DOCX import fixture (tests/library.test.mjs): a Title paragraph (the cover title), then 5 chapters
(Heading 1), each with 2 sections (Heading 2), a paragraph per section and a table in chapter 3. Arabic.

    uv run --no-project --with python-docx python tests/fixtures/import/make_docx5.py <out.docx>
"""
import sys

from docx import Document

d = Document()
d.add_paragraph("تصريح استيراد تجريبي", style="Title")
d.add_paragraph("وثيقة متطلبات الخدمة")
for c, title in enumerate(["المقدمة", "بطاقة الخدمة", "المتطلبات الوظيفية", "قواعد العمل", "قصص المستخدمين"], 1):
    d.add_heading(title, 1)
    for s in (1, 2):
        d.add_heading(f"القسم {s} من {title}", 2)
        d.add_paragraph(f"نص القسم {s} في الفصل {c}.")
    if c == 3:
        t = d.add_table(rows=3, cols=2)
        for i, row in enumerate([["الرمز", "المتطلب"], ["FR-01", "تقديم الطلب"], ["FR-02", "دفع الرسوم"]]):
            for j, v in enumerate(row):
                t.cell(i, j).text = v
d.save(sys.argv[1])
